#!/usr/bin/env python3
"""
Extract the hidden Python script from ESP32_ADXL345_3Axis_Telemetry.xlsx

  1. The 3 sheets are the L*, a*, b* planes of an image (CIELAB, D65, 2 deg).
  2. Convert Lab -> sRGB 8-bit (values come out as exact integers).
  3. Flatten row-major as R,G,B,R,G,B,... and read the LSB at each position,
     stepping forward by the Gijswijt sequence. The sequence restarts every
     500 bits; the position keeps going.
  4. Bits are MSB-first. First 8 bytes = big-endian payload length.

Usage: python3 extract.py ESP32_ADXL345_3Axis_Telemetry.xlsx [out.py]
"""
import sys
import numpy as np
from openpyxl import load_workbook


def lab_to_srgb8(L, a, b):
    # CIELAB -> XYZ (D65 white point, 2 degree observer)
    fy = (L + 16.0) / 116.0
    fx = fy + a / 500.0
    fz = fy - b / 200.0
    eps = 6.0 / 29.0

    def finv(t):
        return np.where(t > eps, t ** 3, 3 * eps ** 2 * (t - 4.0 / 29.0))

    X = 0.95047 * finv(fx)
    Y = 1.00000 * finv(fy)
    Z = 1.08883 * finv(fz)

    # XYZ -> linear sRGB
    M = np.array([[3.24048134, -1.53715152, -0.49853633],
                  [-0.96925495, 1.87599, 0.04155593],
                  [0.05564664, -0.20404134, 1.05731107]])
    rgb = np.stack([X, Y, Z], -1) @ M.T

    # linear -> gamma-encoded sRGB
    rgb = np.where(rgb > 0.0031308,
                   1.055 * np.power(np.clip(rgb, 0, None), 1 / 2.4) - 0.055,
                   12.92 * rgb)
    return np.clip(np.round(rgb * 255), 0, 255).astype(np.uint8)


def gijswijt(n_terms):
    """b(n+1) = largest k such that the sequence so far ends in some block repeated k times."""
    seq = [1]
    while len(seq) < n_terms:
        n = len(seq)
        best = 1
        L = 1
        while (best + 1) * L <= n:          # only lengths that could beat 'best'
            k = 1
            tail = seq[n - L:]
            while (k + 1) * L <= n and seq[n - (k + 1) * L:n - k * L] == tail:
                k += 1
            best = max(best, k)
            L += 1
        seq.append(best)
    return seq


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    xlsx = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "epstein_secret.py"

    print("[*] Reading the three Lab sheets...")
    wb = load_workbook(xlsx, read_only=True)
    planes = [np.array(list(ws.iter_rows(values_only=True)), dtype=float) for ws in wb]
    print(f"    sheets: {len(planes)}, shape: {planes[0].shape}")

    print("[*] Converting Lab -> sRGB...")
    img = lab_to_srgb8(*planes)                 # (400, 302, 3)
    lsb = (img.reshape(-1) & 1).astype(np.uint8)

    print("[*] Generating Gijswijt sequence (500 terms)...")
    g = gijswijt(500)

    print("[*] Reading LSBs with Gijswijt stepping (reset every 500 bits)...")
    bits, pos, n = [], 0, 0
    while pos < len(lsb):
        bits.append(lsb[pos])
        pos += g[n % 500]
        n += 1

    data = np.packbits(np.array(bits, dtype=np.uint8)).tobytes()
    length = int.from_bytes(data[:8], "big")
    payload = data[8:8 + length]
    print(f"[+] Payload length: {length} bytes")

    with open(out, "wb") as fh:
        fh.write(payload)
    print(f"[+] Written to {out}\n")
    print(payload.decode("utf-8", "replace"))


if __name__ == "__main__":
    main()
